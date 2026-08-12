"""Penguraian dokumen: PDF ilmiah, pedoman, dan transkrip.

Jurnal dipecah menjadi bagian dan potongan yang menyimpan nomor halaman, agar
jawaban fitur Tanya Jurnal selalu bisa menunjuk halaman sumbernya.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path


@dataclass
class Page:
    number: int
    text: str


@dataclass
class ParsedDocument:
    pages: list[Page]
    title: str = ""

    @property
    def text(self) -> str:
        return "\n\n".join(p.text for p in self.pages)

    @property
    def page_count(self) -> int:
        return len(self.pages)


class PdfSupportMissing(RuntimeError):
    pass


def parse_pdf(path: Path | str) -> ParsedDocument:
    """Baca PDF menjadi halaman-halaman teks.

    Untuk PDF hasil pindaian yang tidak memuat lapisan teks, halaman akan
    kosong; pemanggil bertanggung jawab menyalurkannya ke OCR (Bagian 7.1).
    """
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - bergantung lingkungan
        raise PdfSupportMissing(
            "pdfplumber belum terpasang. Jalankan: pip install pdfplumber"
        ) from exc

    pages: list[Page] = []
    title = ""
    with pdfplumber.open(str(path)) as pdf:
        meta = pdf.metadata or {}
        title = str(meta.get("Title") or "").strip()
        for index, page in enumerate(pdf.pages, start=1):
            pages.append(Page(number=index, text=(page.extract_text() or "").strip()))
    return ParsedDocument(pages=pages, title=title or Path(path).stem)


def parse_text_file(path: Path | str) -> ParsedDocument:
    raw = Path(path).read_text(encoding="utf-8", errors="replace")
    return ParsedDocument(pages=[Page(number=1, text=raw)], title=Path(path).stem)


def parse_document(path: Path | str) -> ParsedDocument:
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix in (".txt", ".md", ".text"):
        return parse_text_file(path)
    if suffix == ".docx":
        return parse_docx(path)
    raise ValueError(f"Format dokumen '{suffix}' belum didukung.")


def parse_docx(path: Path | str) -> ParsedDocument:
    from docx import Document  # type: ignore

    document = Document(str(path))
    text = "\n".join(p.text for p in document.paragraphs)
    return ParsedDocument(pages=[Page(number=1, text=text)], title=Path(path).stem)


def chunk_pages(
    pages: list[Page], target_chars: int = 1200, overlap: int = 150
) -> list[dict]:
    """Potong dokumen menjadi bagian yang tetap membawa nomor halamannya."""
    chunks: list[dict] = []
    ordinal = 0
    for page in pages:
        text = re.sub(r"[ \t]+", " ", page.text).strip()
        if not text:
            continue
        paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
        buffer = ""
        for paragraph in paragraphs:
            if len(buffer) + len(paragraph) + 1 <= target_chars:
                buffer = f"{buffer}\n{paragraph}".strip()
                continue
            if buffer:
                chunks.append({"page": page.number, "ordinal": ordinal, "text": buffer})
                ordinal += 1
                tail = buffer[-overlap:] if overlap else ""
                buffer = f"{tail}\n{paragraph}".strip() if tail else paragraph
            else:
                for start in range(0, len(paragraph), target_chars):
                    chunks.append(
                        {
                            "page": page.number,
                            "ordinal": ordinal,
                            "text": paragraph[start : start + target_chars],
                        }
                    )
                    ordinal += 1
                buffer = ""
        if buffer:
            chunks.append({"page": page.number, "ordinal": ordinal, "text": buffer})
            ordinal += 1
    return chunks


#: Bagian baku artikel ilmiah, dipakai memecah jurnal jadi bagian bernama.
SECTION_HEADINGS = (
    "abstract",
    "abstrak",
    "introduction",
    "pendahuluan",
    "literature review",
    "tinjauan pustaka",
    "method",
    "methods",
    "methodology",
    "metode",
    "metodologi",
    "result",
    "results",
    "hasil",
    "discussion",
    "pembahasan",
    "conclusion",
    "simpulan",
    "kesimpulan",
    "references",
    "daftar pustaka",
)


def split_sections(text: str) -> dict[str, str]:
    """Pecah artikel menjadi bagian bernama berdasarkan judul baku."""
    lines = text.splitlines()
    sections: dict[str, list[str]] = {}
    current = "awal"
    for line in lines:
        stripped = line.strip()
        normalized = re.sub(r"^[0-9.\s]+", "", stripped).strip().lower()
        if normalized in SECTION_HEADINGS and len(stripped) < 60:
            current = normalized
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    return {key: "\n".join(value).strip() for key, value in sections.items() if any(value)}
